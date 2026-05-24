"""XLSX writer — long-form (one row per cameo) per §6.3."""

from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from cameo_convert.model import InvertedDataset
from cameo_convert.writers.base import FLAT_COLUMNS, flat_rows, metadata_dict


class XlsxWriter:
    name = "xlsx"
    extension = "xlsx"

    def write(self, dataset: InvertedDataset, path: Path) -> None:
        wb = openpyxl.Workbook()

        # metadata sheet
        meta_ws = wb.active
        meta_ws.title = "metadata"
        for row in metadata_dict(dataset).items():
            meta_ws.append([row[0], str(row[1]) if row[1] is not None else ""])
        for cell in meta_ws["A"]:
            cell.font = Font(bold=True)

        ws = wb.create_sheet("cameos")
        ws.append(FLAT_COLUMNS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in flat_rows(dataset):
            ws.append([row[col] for col in FLAT_COLUMNS])

        ws.freeze_panes = "A2"
        for i, col in enumerate(FLAT_COLUMNS, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(
                12, min(40, len(col) + 4)
            )

        wb.save(path)
