import json
import sqlite3
import zipfile
from pathlib import Path

import openpyxl
import pytest
import tomllib
import yaml

from cameo_convert.filter import clean, full
from cameo_convert.reader import read
from cameo_convert.sets import load as load_set_dates
from cameo_convert.transform import invert
from cameo_convert.writers import ALL_FORMATS, REGISTRY

FIXTURE = Path(__file__).parent / "fixtures" / "tiny_sample.ods"
SET_DATES_FIXTURE = Path(__file__).parent / "fixtures" / "set_release_dates.json"


@pytest.fixture(scope="module")
def dataset():
    entries, meta = read(FIXTURE)
    return invert(
        entries,
        source_file="tests/fixtures/tiny_sample.ods",
        source_last_updated=meta.last_updated,
        release_tag="v0.test",
        set_date_resolver=load_set_dates(SET_DATES_FIXTURE),
    )


@pytest.mark.parametrize("fmt", ALL_FORMATS)
def test_writer_produces_file(dataset, tmp_path, fmt):
    writer = REGISTRY[fmt]()
    path = tmp_path / f"out.{writer.extension}"
    writer.write(full(dataset), path)
    assert path.exists()
    assert path.stat().st_size > 0


def test_json_schema_matches_design_doc(dataset, tmp_path):
    path = tmp_path / "out.json"
    REGISTRY["json"]().write(full(dataset), path)
    obj = json.loads(path.read_text(encoding="utf-8"))
    assert "metadata" in obj and "cards" in obj
    meta = obj["metadata"]
    for k in (
        "source_file",
        "source_last_updated",
        "generated_at",
        "release_tag",
        "variant",
        "total_cards",
        "total_cameo_entries",
    ):
        assert k in meta
    # §6.4 worked example: Aquapolis|Town Volunteers|136 contains Bulbasaur.
    assert "Aquapolis|Town Volunteers|136" in obj["cards"]
    card = obj["cards"]["Aquapolis|Town Volunteers|136"]
    assert any(c["subject"] == "Bulbasaur" for c in card["cameos"])
    # Different-form preserved in structured output (D7).
    mega = obj["cards"]["Made Up Set|Some Card|1"]
    cam = mega["cameos"][0]
    assert cam["subject"] == "Mega Bulbasaur"
    assert cam["parent_species"] == "Bulbasaur"
    assert "different_form" in cam["flags"]


def test_csv_d7_double_rows_for_different_form(dataset, tmp_path):
    path = tmp_path / "out.csv"
    REGISTRY["csv"]().write(full(dataset), path)
    rows = path.read_text(encoding="utf-8").splitlines()
    # Find rows for Mega Bulbasaur's card
    mega_rows = [r for r in rows if "Some Card" in r and "Made Up Set" in r]
    # Expect 2 rows: one literal subject, one species_parent.
    assert len(mega_rows) == 2
    assert any("literal" in r and "Mega Bulbasaur" in r for r in mega_rows)
    assert any("species_parent" in r and "Bulbasaur" in r for r in mega_rows)


def test_sqlite_two_tables_and_counts(dataset, tmp_path):
    path = tmp_path / "out.sqlite"
    REGISTRY["sqlite"]().write(full(dataset), path)
    con = sqlite3.connect(path)
    try:
        tables = {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        assert {"cards", "cameos", "metadata"} <= tables
        n_cards = con.execute("SELECT COUNT(*) FROM cards").fetchone()[0]
        assert n_cards == dataset.total_cards
        n_cameos = con.execute("SELECT COUNT(*) FROM cameos").fetchone()[0]
        assert n_cameos == dataset.total_cameo_entries
    finally:
        con.close()


def test_xlsx_has_cameos_sheet(dataset, tmp_path):
    path = tmp_path / "out.xlsx"
    REGISTRY["xlsx"]().write(full(dataset), path)
    wb = openpyxl.load_workbook(path)
    assert "cameos" in wb.sheetnames
    assert "metadata" in wb.sheetnames


def test_ods_is_valid_zip_with_content(dataset, tmp_path):
    path = tmp_path / "out.ods"
    REGISTRY["ods"]().write(full(dataset), path)
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        assert {"mimetype", "META-INF/manifest.xml", "meta.xml", "content.xml"} <= names
        assert z.read("mimetype") == b"application/vnd.oasis.opendocument.spreadsheet"


def test_yaml_roundtrip_via_dict(dataset, tmp_path):
    path = tmp_path / "out.yaml"
    REGISTRY["yaml"]().write(full(dataset), path)
    obj = yaml.safe_load(path.read_text(encoding="utf-8"))
    assert obj["metadata"]["total_cards"] == dataset.total_cards


def test_toml_roundtrip_via_dict(dataset, tmp_path):
    path = tmp_path / "out.toml"
    REGISTRY["toml"]().write(full(dataset), path)
    obj = tomllib.loads(path.read_text(encoding="utf-8"))
    assert obj["metadata"]["total_cards"] == dataset.total_cards
    assert len(obj["cards"]) == dataset.total_cards


def test_clean_variant_drops_excluded(dataset, tmp_path):
    cleaned = clean(dataset)
    # Blank-card-name rows always drop.
    for k in cleaned.cards:
        assert cleaned.cards[k].card_name is not None
    # Standalone non-English/italic cards drop, but reprint-siblings of an
    # English card are kept (with their flags preserved). So we don't assert
    # the absence of NON_ENGLISH/ITALIC flags here — the survivors are
    # specifically the ones whose artwork-group has a clean sibling.
    path = tmp_path / "out.json"
    REGISTRY["json"]().write(cleaned, path)
    obj = json.loads(path.read_text(encoding="utf-8"))
    assert obj["metadata"]["variant"] == "clean"
    # The standalone JP-only card is gone.
    assert "Japanese Set|JP-Only Card|JP1" not in obj["cards"]
    # The JP reprint of an English card survives (sibling rule).
    assert "Japanese Set|Reprint Test|JP5" in obj["cards"]


def test_json_includes_primary_pokemons(dataset, tmp_path):
    path = tmp_path / "out.json"
    REGISTRY["json"]().write(full(dataset), path)
    obj = json.loads(path.read_text(encoding="utf-8"))
    dual = obj["cards"]["Cosmic Eclipse|Bulbasaur & Ivysaur-GX|999"]
    assert dual["primary_pokemons"] == ["Bulbasaur", "Ivysaur"]
    red = obj["cards"]["SM-P Promos|Red's Pikachu|270"]
    assert red["primary_pokemons"] == ["Pikachu"]


def test_md_groups_by_primary_pokemon(dataset, tmp_path):
    path = tmp_path / "out.md"
    REGISTRY["md"]().write(full(dataset), path)
    text = path.read_text(encoding="utf-8")
    # Outer headers (lines starting with "## ") are Pokémon names, not sets.
    h2s = [line[3:].strip() for line in text.splitlines() if line.startswith("## ")]
    # Multi-name card should make both Bulbasaur and Ivysaur appear.
    assert "Bulbasaur" in h2s
    assert "Ivysaur" in h2s
    # The dual-named card appears under both — section is from the
    # "## Bulbasaur\n" header line to the next "\n## " header.
    bulb_section = text.split("## Bulbasaur\n")[1].split("\n## ")[0]
    assert "Bulbasaur & Ivysaur-GX" in bulb_section
    ivy_section = text.split("## Ivysaur\n")[1].split("\n## ")[0]
    assert "Bulbasaur & Ivysaur-GX" in ivy_section


def test_sqlite_card_primary_pokemons_join_table(dataset, tmp_path):
    path = tmp_path / "out.sqlite"
    REGISTRY["sqlite"]().write(full(dataset), path)
    con = sqlite3.connect(path)
    try:
        tables = {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        assert "card_primary_pokemons" in tables
        # Bulbasaur should be a primary for the dual-name card.
        cards_for_bulbasaur = [
            r[0] for r in con.execute(
                "SELECT card_key FROM card_primary_pokemons WHERE pokemon = 'Bulbasaur'"
            )
        ]
        assert "Cosmic Eclipse|Bulbasaur & Ivysaur-GX|999" in cards_for_bulbasaur
    finally:
        con.close()


def test_csv_has_primary_pokemons_column(dataset, tmp_path):
    path = tmp_path / "out.csv"
    REGISTRY["csv"]().write(full(dataset), path)
    rows = path.read_text(encoding="utf-8").splitlines()
    assert "primary_pokemons" in rows[0]
    dual_row = next(r for r in rows if "Bulbasaur & Ivysaur-GX" in r)
    assert "Bulbasaur,Ivysaur" in dual_row


def test_json_includes_release_date(dataset, tmp_path):
    path = tmp_path / "out.json"
    REGISTRY["json"]().write(full(dataset), path)
    obj = json.loads(path.read_text(encoding="utf-8"))
    # Aquapolis is in our fixture set_release_dates.json → 2003-01-15
    aqua = obj["cards"]["Aquapolis|Town Volunteers|136"]
    assert aqua["release_date"] == "2003-01-15"
    # "Made Up Set" is not in the fixture → null
    made_up = obj["cards"]["Made Up Set|Some Card|1"]
    assert made_up["release_date"] is None


def test_md_sorts_cards_chronologically_within_bucket(dataset, tmp_path):
    # Build a hand-rolled dataset with three cards in known dates so the
    # ordering is unambiguous. Use "Eevee" as both cameo subject AND card
    # name so card_primary_pokemons resolves to ["Eevee"] — the cards bucket
    # under ## Eevee.
    from cameo_convert.model import CameoEntry

    entries = [
        CameoEntry(
            cameo_subject="Eevee",
            parent_species=None,
            subject_kind="pokemon",
            ndex=133,
            region=None,
            card_name="Eevee",
            set_name=set_name,
            collector_number="1",
            notes=None,
            flags=frozenset(),
            artwork_group_id=None,
        )
        for set_name in ("OLD SET", "MID SET", "NEW SET")
    ]
    dates = {
        "OLD SET": "2000-01-01",
        "MID SET": "2010-01-01",
        "NEW SET": "2020-01-01",
    }
    resolver = lambda s: dates.get(s)  # noqa: E731
    ds = invert(entries, source_file="x", set_date_resolver=resolver)
    path = tmp_path / "out.md"
    REGISTRY["md"]().write(full(ds), path)
    text = path.read_text(encoding="utf-8")
    # All three Eevee cards appear under ## Eevee — verify chronological order.
    eevee = text.split("## Eevee\n", 1)[1].split("\n## ")[0]
    old_pos = eevee.find("OLD SET")
    mid_pos = eevee.find("MID SET")
    new_pos = eevee.find("NEW SET")
    assert 0 <= old_pos < mid_pos < new_pos, (
        f"expected chronological order OLD < MID < NEW, got {old_pos}/{mid_pos}/{new_pos}"
    )


def test_sqlite_release_date_column(dataset, tmp_path):
    path = tmp_path / "out.sqlite"
    REGISTRY["sqlite"]().write(full(dataset), path)
    con = sqlite3.connect(path)
    try:
        rows = con.execute(
            "SELECT release_date FROM cards WHERE set_name = 'Aquapolis'"
        ).fetchall()
        assert rows
        for (rd,) in rows:
            assert rd == "2003-01-15"
    finally:
        con.close()
